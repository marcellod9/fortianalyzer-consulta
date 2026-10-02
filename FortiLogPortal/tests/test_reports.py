"""Relatórios de acesso: contagens, gráficos, amostra e exportação (dados de demonstração)."""
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services import report_export, reports

END = datetime(2026, 10, 1, 18, 0)


def q(**kw):
    return reports.ReportQuery(start=END - timedelta(hours=24), end=END, **kw)


def test_value_is_required_except_for_general_report():
    with pytest.raises(ValueError, match="usuário"):
        q(tipo="usuario")
    with pytest.raises(ValueError):
        q(tipo="ip", valor="não é ip")
    assert q(tipo="geral", valor="  ").valor is None
    assert q(tipo="usuario", valor="MARISTA\\maria").base_query("webfilter").filter_expr() == 'user~"maria"'
    assert q(tipo="site", valor="https://www.facebook.com/x").base_query("webfilter").filter_expr() == 'hostname~"facebook.com"'
    assert q(tipo="aplicacao", valor="YouTube", somente_bloqueios=True).base_query("app-ctrl").filter_expr() == \
        'app~"YouTube" and action=block'


def test_general_report_counts_and_charts():
    r = reports.run_report(q(tipo="geral"), use_cache=False)
    assert set(r["fontes"]) == {"webfilter", "app-ctrl"} and r["erros"] == {}
    s = r["resumo"]
    assert s["eventos"] == sum(f["total"] for f in r["fontes"].values()) and s["permitidos"] + s["bloqueados"] == s["lidos"]
    by_id = {c["id"]: c for c in r["graficos"]}
    assert {"acao", "categorias", "sites", "aplicacoes", "usuarios", "tempo"} <= set(by_id)
    assert sum(n for _, n in by_id["acao"]["itens"]) == s["lidos"]
    # pizza: no máximo 7 fatias + "Outros", e a soma fecha com a base
    cat = by_id["categorias"]
    assert len(cat["itens"]) <= 8 and sum(n for _, n in cat["itens"]) + cat["sem_valor"] == cat["base"]
    assert all(c["fonte"] == "filtro web" for c in (by_id["sites"], cat))
    t = by_id["tempo"]
    assert t["agrupamento"] == "hora" and len(t["rotulos"]) == 25 and sum(t["permitidos"]) + sum(t["bloqueados"]) == s["lidos"]


def test_report_by_site_uses_webfilter_only_and_drops_its_own_dimension():
    r = reports.run_report(q(tipo="site", valor="facebook.com"), use_cache=False)
    assert list(r["fontes"]) == ["webfilter"]
    ids = [c["id"] for c in r["graficos"]]
    assert "sites" not in ids and "usuarios" in ids
    assert all(e["site"].endswith("facebook.com") for e in r["eventos"])


def test_sample_warning_when_period_has_more_events(monkeypatch):
    monkeypatch.setattr(reports.settings, "report_max_rows", 200)
    r = reports.run_report(q(tipo="geral"), use_cache=False)
    assert r["amostra"] and "200 de" in r["aviso_amostra"] and r["resumo"]["lidos"] == 400
    tempo = next(c for c in r["graficos"] if c["id"] == "tempo")
    assert len(tempo["rotulos"]) < 25  # a linha do tempo começa no evento mais antigo lido


def test_long_period_groups_by_day():
    r = reports.run_report(reports.ReportQuery(start=END - timedelta(days=7), end=END, tipo="geral"), use_cache=False)
    t = next(c for c in r["graficos"] if c["id"] == "tempo")
    assert t["agrupamento"] == "dia" and len(t["rotulos"]) == 8


def test_pdf_and_excel_exports():
    r = reports.run_report(q(tipo="usuario", valor="maria"), use_cache=False)
    pdf, name = report_export.build("pdf", r, reports.EVENT_COLUMNS)
    assert pdf.startswith(b"%PDF") and name.endswith(".pdf")
    xlsx, _ = report_export.build("xlsx", r, reports.EVENT_COLUMNS)
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames[0] == "Resumo" and wb.sheetnames[-1] == "Eventos"
    assert wb["Eventos"].max_row == len(r["eventos"]) + 1
    assert any(ws._charts for ws in wb.worksheets)
    assert report_export.time_label("2026-10-01 18", "hora") == "01/10 18h"
    assert report_export._step(7) == 2 and report_export._step(1234) == 500


def test_report_api_download_and_history():
    with TestClient(app) as c:
        r = c.post("/api/reports", json={"tipo": "ip", "valor": "10.10.0.0/16", "start": "2026-10-01T00:00", "end": "2026-10-01T18:00"},
                   headers={"X-Portal-User": "ana"})
        assert r.status_code == 200, r.text
        d = r.json()
        assert "eventos" not in d and d["eventos_exportaveis"] == d["resumo"]["lidos"]
        pdf = c.get(f"/api/reports/{d['report_id']}.pdf")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        assert c.get(f"/api/reports/{d['report_id']}.csv").status_code == 400
        assert c.get("/api/reports/naoexiste.pdf").status_code == 404
        bad = c.post("/api/reports", json={"tipo": "usuario", "start": "2026-10-01T00:00", "end": "2026-10-01T18:00"})
        assert bad.status_code in (400, 422) and "usuário" in bad.text
        hist = c.get("/api/history?type=relatorio").json()
        assert any("10.10.0.0/16" in h["term"] for h in hist)
        assert c.get("/relatorios").status_code == 200


def test_blocked_only_report_renames_and_skips_redundant_charts():
    r = reports.run_report(q(tipo="geral", somente_bloqueios=True), use_cache=False)
    titles = {c["id"]: c["titulo"] for c in r["graficos"]}
    assert "acao" not in titles and "sites_bloqueados" not in titles
    assert titles["sites"] == "Sites mais bloqueados" and titles["tempo"] == "Bloqueios ao longo do tempo"
    assert r["resumo"]["permitidos"] == 0


def test_site_report_lists_every_user_with_counts_and_dates():
    r = reports.run_report(q(tipo="site", valor="facebook.com"), use_cache=False)
    d = r["detalhe"]
    assert d["agrupamento"] == "usuario" and d["completo"] and d["titulo"] == "Usuários que acessaram facebook.com"
    assert sum(x["acessos"] for x in d["linhas"]) == r["resumo"]["lidos"]
    users = {e["usuario"] or f"(sem login) {e['ip_origem']}" for e in r["eventos"]}
    assert {x["usuario"] for x in d["linhas"]} == users
    top = d["linhas"][0]
    assert top["primeiro"] <= top["ultimo"] and top["permitidos"] + top["bloqueados"] == top["acessos"]
    assert [c for c, _ in d["colunas"]][:3] == ["usuario", "ips", "maquinas"]


def test_user_report_lists_destinations():
    d = reports.run_report(q(tipo="usuario", valor="maria"), use_cache=False)["detalhe"]
    assert d["agrupamento"] == "destino" and {x["tipo"] for x in d["linhas"]} == {"Site", "Aplicação"}


def _ev(i, user, blocked=False):
    return {"data_hora": f"2026-10-01 {10 + i // 60:02d}:{i % 60:02d}:00", "usuario": user, "ip_origem": f"10.0.0.{i % 250}",
            "maquina": "", "firewall": "FW-A", "site": "facebook.com", "ip_destino": "", "aplicacao": "", "categoria": "Social",
            "bloqueado": blocked, "tipo_log": "webfilter", "situacao": "bloqueado" if blocked else "permitido", "regra": "1"}


def test_extra_searches_complete_the_user_list_when_the_period_has_more_events(monkeypatch):
    """O FAZ tem 30 eventos de ana (mais recentes) e 1 de bia e 1 de caio (mais antigos); o relatório lê só 10."""
    monkeypatch.setattr(reports.settings, "report_max_rows_filtered", 10)
    events = [_ev(100 - i, "ana") for i in range(30)] + [_ev(5, "bia", True), _ev(3, "caio")]
    filters = []

    def fake_run_query(lq, **kw):
        flt = lq.filter_expr()
        filters.append(flt)
        excluded = [part.split('"')[1] for part in flt.split(" and ") if part.startswith("user!=")]
        match = [e for e in events if e["usuario"] not in excluded]
        rows = match[: lq.limit]
        return {"total": len(match), "returned": len(rows), "rows": rows}
    monkeypatch.setattr(reports, "run_query", fake_run_query)
    r = reports.run_report(q(tipo="site", valor="facebook.com"), use_cache=False)
    d = r["detalhe"]
    assert [x["usuario"] for x in d["linhas"]] == ["ana", "bia", "caio"] and d["completo"]
    assert d["linhas"][0]["acessos"] == 10 and not d["linhas"][0]["parcial"]
    assert d["linhas"][1]["parcial"] and d["linhas"][1]["bloqueados"] == 1
    assert 'user!="ana"' in filters[1] and r["amostra"] and "lista de usuários está completa" in r["aviso_amostra"]
    # gráficos continuam só com os eventos principais (as buscas extras não distorcem as contagens)
    assert sum(n for _, n in next(c for c in r["graficos"] if c["id"] == "usuarios")["itens"]) == 10


def test_user_list_is_flagged_incomplete_when_remaining_events_have_no_login(monkeypatch):
    monkeypatch.setattr(reports.settings, "report_max_rows_filtered", 5)
    monkeypatch.setattr(reports, "DISCOVERY_LIMIT", 3)  # sobram 10 eventos sem login: a busca extra não lê todos
    events = [_ev(50 - i, "ana") for i in range(10)] + [_ev(20 - i, "") for i in range(10)]

    def fake_run_query(lq, **kw):
        excluded = [p.split('"')[1] for p in lq.filter_expr().split(" and ") if p.startswith("user!=")]
        match = [e for e in events if e["usuario"] not in excluded]
        return {"total": len(match), "returned": min(len(match), lq.limit), "rows": match[: lq.limit]}
    monkeypatch.setattr(reports, "run_query", fake_run_query)
    r = reports.run_report(q(tipo="site", valor="facebook.com"), use_cache=False)
    assert not r["detalhe"]["completo"] and "pode não estar completa" in r["aviso_amostra"]


def test_detail_in_pdf_and_excel():
    r = reports.run_report(q(tipo="site", valor="facebook.com"), use_cache=False)
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(report_export.build("xlsx", r, reports.EVENT_COLUMNS)[0]))
    assert wb.sheetnames[1] == "Usuários"
    ws = wb["Usuários"]
    assert ws.cell(3, 1).value == "Usuário" and ws.max_row == 3 + len(r["detalhe"]["linhas"])
    assert report_export.detail_rows(r["detalhe"])[0][6].count("/") == 2
    assert report_export.build("pdf", r, reports.EVENT_COLUMNS)[0].startswith(b"%PDF")
