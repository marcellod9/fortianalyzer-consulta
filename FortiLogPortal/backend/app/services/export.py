"""Exportação de resultados para CSV, XLSX e PDF (cópia salva em exports/)."""
import csv
import io
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from ..config import settings

FORMATS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, (list, tuple)):
        return ", ".join(_cell(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_cell(x)}" for k, x in v.items())
    s = str(v)
    # evita injeção de fórmula ao abrir no Excel
    if s[:1] in ("=", "+", "-", "@") and not re.fullmatch(r"-?\d+(\.\d+)?", s):
        s = "'" + s
    return s


def to_csv(columns, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)  # ";" abre direto no Excel pt-BR
    w.writerow([label for _, label in columns])
    for r in rows:
        w.writerow([_cell(r.get(k)) for k, _ in columns])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(columns, rows, title: str, meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Gerado em {datetime.now():%d/%m/%Y %H:%M} pelo FortiLogPortal"] +
              [f"{k}: {_cell(v)}" for k, v in (meta or {}).items()])
    ws.append([])
    ws.append([label for _, label in columns])
    head = ws.max_row
    fill = PatternFill("solid", fgColor="1F4E79")
    for c in ws[head]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_cell(r.get(k)) for k, _ in columns])
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(columns))}{ws.max_row}"
    for i, (k, label) in enumerate(columns, start=1):
        width = max([len(label)] + [len(_cell(r.get(k))) for r in rows[:300]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(columns, rows, title: str, meta: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from xml.sax.saxutils import escape

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=6.5, leading=8)
    headst = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    story = [Paragraph(escape(title), styles["Title"]),
             Paragraph(escape(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} pelo FortiLogPortal. " +
                              " | ".join(f"{k}: {_cell(v)}" for k, v in (meta or {}).items())), styles["Normal"]),
             Spacer(1, 4 * mm)]
    data = [[Paragraph(escape(label), headst) for _, label in columns]]
    for r in rows:
        data.append([Paragraph(escape(_cell(r.get(k))[:600]), cell) for k, _ in columns])
    width = doc.width / max(len(columns), 1)
    t = LongTable(data, colWidths=[width] * len(columns), repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t)
    if not rows:
        story.append(Paragraph("Nenhum registro.", styles["Normal"]))
    doc.build(story)
    return out.getvalue()


def build(fmt: str, title: str, columns, rows, meta: dict | None = None) -> tuple[bytes, str]:
    if fmt not in FORMATS:
        raise ValueError("Formato deve ser csv, xlsx ou pdf")
    columns = [tuple(c) for c in columns]
    if fmt == "csv":
        content = to_csv(columns, rows)
    elif fmt == "xlsx":
        content = to_xlsx(columns, rows, title, meta or {})
    else:
        content = to_pdf(columns, rows, title, meta or {})
    ascii_title = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9\-]+", "_", ascii_title.lower()).strip("_")[:40] or "resultado"
    filename = f"{slug}_{datetime.now():%Y%m%d_%H%M%S}.{fmt}"
    settings.export_dir.mkdir(parents=True, exist_ok=True)
    Path(settings.export_dir / filename).write_bytes(content)
    return content, filename
