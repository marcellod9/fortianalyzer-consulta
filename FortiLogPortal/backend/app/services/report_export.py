"""Relatório de acessos em PDF (com gráficos, no estilo dos relatórios do FortiAnalyzer) e em Excel."""
import io
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from ..config import settings
from .export import FORMATS, _cell

# Paleta categórica (ordem fixa, validada para daltonismo) e cores de situação, as mesmas da tela
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
OTHERS = "#a3a29c"
GOOD, CRITICAL = "#0ca30c", "#d03b3b"
HEAD = "#1F4E79"


def _pct(n: int, total: int) -> str:
    return f"{(100 * n / total):.1f}%".replace(".", ",") if total else "-"


def _num(n) -> str:
    return f"{n:,}".replace(",", ".") if isinstance(n, int) else str(n)


def time_label(key: str, group: str) -> str:
    """'2026-10-01 18' -> '01/10 18h'; '2026-10-01' -> '01/10'."""
    d = key[8:10] + "/" + key[5:7]
    return f"{d} {key[11:13]}h" if group == "hora" else d


def _meta(rep: dict) -> list[tuple[str, str]]:
    r = rep["resumo"]
    out = [("Período", f"{rep['periodo']['inicio']} a {rep['periodo']['fim']}"), ("Firewall", rep.get("firewall") or "Todos"),
           ("Eventos no FortiAnalyzer", _num(r["eventos"])), ("Eventos analisados", _num(r["lidos"])),
           ("Permitidos", _num(r["permitidos"])), ("Bloqueados", _num(r["bloqueados"])),
           ("Usuários", _num(r["usuarios"])), ("IPs de origem", _num(r["ips"])),
           ("Sites", _num(r["sites"])), ("Aplicações", _num(r["aplicacoes"]))]
    return out


def _notes(rep: dict) -> list[str]:
    return [n for n in (rep.get("aviso_amostra"), rep.get("aviso")) if n]


def _rows(chart: dict) -> list[list[str]]:
    """Linhas da tabela de cada gráfico: item, registros e % (ou horário, permitidos e bloqueados)."""
    if chart["tipo"] == "line":
        return [[time_label(k, chart["agrupamento"]), _num(a), _num(b)] for k, a, b in chart["itens"]]
    total = sum(n for _, n in chart["itens"]) if chart["tipo"] == "pie" or chart["id"] == "acao" else chart["base"]
    return [[str(k), _num(n), _pct(n, total)] for k, n in chart["itens"]]


def _headers(chart: dict) -> list[str]:
    if chart["tipo"] == "line":
        return ["Horário" if chart["agrupamento"] == "hora" else "Dia", "Permitidos", "Bloqueados"]
    return [chart["coluna"], "Registros", "%"]


def detail_rows(det: dict) -> list[list[str]]:
    """Linhas da tabela detalhada (usuários ou destinos); * marca contagem parcial (busca extra)."""
    out = []
    for line in det["linhas"]:
        row = []
        for k, _ in det["colunas"]:
            v = line.get(k, "")
            if k in ("primeiro", "ultimo"):
                v = f"{v[8:10]}/{v[5:7]}/{v[:4]} {v[11:16]}" if v else ""
            elif isinstance(v, int):
                v = _num(v) + ("*" if k == "acessos" and line.get("parcial") else "")
            row.append(str(v))
        out.append(row)
    return out


def detail_note(det: dict) -> str:
    if det["agrupamento"] == "destino":
        n = f"{_num(det['total'])} destinos."
    else:
        n = f"{_num(det['usuarios'])} usuário(s)"
        n += f" e {_num(det['sem_login'])} IP(s) sem login (no fim da lista)." if det["sem_login"] else "."
    if det["total"] > len(det["linhas"]):
        n += f" Mostrando os {_num(len(det['linhas']))} com mais acessos."
    if det["parciais"]:
        n += " * Encontrado numa busca extra: a contagem desse usuário é parcial."
    if not det["completo"]:
        n += " A lista pode não estar completa (veja o aviso acima)."
    return n


def _colors(chart: dict) -> list[str]:
    if chart["id"] == "acao":
        return [GOOD if k == "Permitidos" else CRITICAL for k, _ in chart["itens"]]
    return [OTHERS if k == "Outros" else SERIES[i % len(SERIES)] for i, (k, _) in enumerate(chart["itens"])]


def _step(top: int) -> int:
    """Intervalo inteiro e "redondo" do eixo (1, 2, 5, 10, 20, 50...) para uns 5 traços."""
    raw, base = max(top, 1) / 5, 1
    while base * 10 <= raw:
        base *= 10
    return next(m * base for m in (1, 2, 5, 10) if m * base >= raw)


# ---- PDF -------------------------------------------------------------------------------------
def to_pdf(rep: dict) -> bytes:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    hx = colors.HexColor
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=rep["titulo"])
    st = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=st["Normal"], fontSize=8, leading=10)
    cell = ParagraphStyle("cell", parent=small, fontSize=7.5, leading=9)
    headc = ParagraphStyle("headc", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    h2 = ParagraphStyle("h2", parent=st["Heading2"], fontSize=12, spaceBefore=6, spaceAfter=4, textColor=hx(HEAD))

    def table(head, rows, widths):
        data = [[Paragraph(escape(h), headc) for h in head]] + [[Paragraph(escape(v[:120]), cell) for v in r] for r in rows]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), hx(HEAD)), ("GRID", (0, 0), (-1, -1), 0.25, hx("#BBBBBB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hx("#F2F6FA")]), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ]))
        return t

    def short(s: str, n: int) -> str:
        return s if len(s) <= n else s[: n - 1] + "…"

    def pie(chart):
        d = Drawing(85 * mm, 62 * mm)
        p = Pie()
        p.x, p.y, p.width, p.height = 4 * mm, 6 * mm, 50 * mm, 50 * mm
        p.data = [n for _, n in chart["itens"]]
        p.labels = None
        p.simpleLabels = 1
        p.slices.strokeColor = colors.white
        p.slices.strokeWidth = 1.5
        for i, c in enumerate(_colors(chart)):
            p.slices[i].fillColor = hx(c)
        p.innerRadiusFraction = 0.55  # rosca: o centro mostra o total
        d.add(p)
        total = sum(p.data)
        d.add(String(29 * mm, 30 * mm, _num(total), textAnchor="middle", fontName="Helvetica-Bold", fontSize=10))
        d.add(String(29 * mm, 26 * mm, "registros", textAnchor="middle", fontName="Helvetica", fontSize=6.5, fillColor=hx("#52514e")))
        lg = Legend()
        lg.x, lg.y = 58 * mm, 54 * mm
        lg.fontSize = 6.5
        lg.dy, lg.dx = 6, 6
        lg.deltay = 8
        lg.alignment = "right"
        lg.columnMaximum = 10
        lg.fontName = "Helvetica"
        lg.colorNamePairs = [(hx(c), f"{short(str(k), 17)} ({_pct(n, total)})") for c, (k, n) in zip(_colors(chart), chart["itens"])]
        d.add(lg)
        return d

    def hbar(chart):
        items = chart["itens"]
        h = max(30, 7 * len(items) + 12) * mm / 1.6
        d = Drawing(85 * mm, h)
        b = HorizontalBarChart()
        b.x, b.y, b.width, b.height = 38 * mm, 4 * mm, 44 * mm, h - 8 * mm
        b.data = [[n for _, n in reversed(items)]]
        b.categoryAxis.categoryNames = [short(str(k), 26) for k, _ in reversed(items)]
        b.categoryAxis.labels.fontSize = 6.5
        b.categoryAxis.labels.boxAnchor = "e"
        b.categoryAxis.strokeColor = hx("#cccccc")
        b.valueAxis.valueMin = 0
        b.valueAxis.valueStep = _step(max(n for _, n in items))
        b.valueAxis.labels.fontSize = 6
        b.valueAxis.labels.fontName = b.categoryAxis.labels.fontName = "Helvetica"
        b.valueAxis.strokeColor = hx("#cccccc")
        b.valueAxis.visibleGrid = True
        b.valueAxis.gridStrokeColor = hx("#eeeeee")
        b.bars[0].fillColor = hx(SERIES[0])
        b.bars[0].strokeColor = None
        b.barSpacing = 1
        d.add(b)
        return d

    def stacked(chart):
        d = Drawing(182 * mm, 62 * mm)
        b = VerticalBarChart()
        b.x, b.y, b.width, b.height = 12 * mm, 12 * mm, 166 * mm, 44 * mm
        b.data = [chart["permitidos"], chart["bloqueados"]]
        b.categoryAxis.style = "stacked"
        names = [time_label(k, chart["agrupamento"]) for k in chart["rotulos"]]
        every = max(1, len(names) // 12)  # rótulos espaçados para não encavalar
        b.categoryAxis.categoryNames = [n if i % every == 0 else "" for i, n in enumerate(names)]
        b.categoryAxis.labels.fontSize = 6
        b.categoryAxis.labels.angle = 0
        b.categoryAxis.strokeColor = hx("#cccccc")
        b.valueAxis.valueMin = 0
        b.valueAxis.valueStep = _step(max([a + c for a, c in zip(chart["permitidos"], chart["bloqueados"])] or [1]))
        b.valueAxis.labels.fontSize = 6
        b.valueAxis.labels.fontName = b.categoryAxis.labels.fontName = "Helvetica"
        b.valueAxis.visibleGrid = True
        b.valueAxis.gridStrokeColor = hx("#eeeeee")
        b.valueAxis.strokeColor = hx("#cccccc")
        b.bars[0].fillColor, b.bars[1].fillColor = hx(GOOD), hx(CRITICAL)
        b.bars[0].strokeColor = b.bars[1].strokeColor = colors.white
        b.bars.strokeWidth = 0.5
        d.add(b)
        series = [(c, label) for c, label, vals in ((GOOD, "Permitidos", chart["permitidos"]),
                                                    (CRITICAL, "Bloqueados", chart["bloqueados"])) if any(vals)]
        for i, (c, label) in enumerate(series):
            d.add(Rect(12 * mm + i * 30 * mm, 2 * mm, 3 * mm, 3 * mm, fillColor=hx(c), strokeColor=None))
            d.add(String(16 * mm + i * 30 * mm, 2.4 * mm, label, fontName="Helvetica", fontSize=7))
        return d

    def meter(chart):
        """Permitidos x bloqueados: uma barra 100% (duas fatias leem melhor numa barra que numa pizza)."""
        d = Drawing(182 * mm, 16 * mm)
        total = sum(n for _, n in chart["itens"]) or 1
        x = 0.0
        for (k, n), c in zip(chart["itens"], _colors(chart)):
            w = 182 * mm * n / total
            d.add(Rect(x, 8 * mm, max(w - 1, 0.5), 6 * mm, fillColor=hx(c), strokeColor=None))
            if w > 30 * mm:
                d.add(String(x + 2 * mm, 3 * mm, f"{k}: {_num(n)} ({_pct(n, total)})", fontName="Helvetica", fontSize=8))
            x += w
        return d

    story = [Paragraph(escape(rep["titulo"]), st["Title"]),
             Paragraph(escape(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} pelo FortiLogPortal com logs do FortiAnalyzer"
                              f" (filtro web e controle de aplicações)."), small), Spacer(1, 3 * mm)]
    meta = _meta(rep)
    half = (len(meta) + 1) // 2
    kv = [[Paragraph(f"<b>{escape(a)}</b>", small), Paragraph(escape(b), small),
           Paragraph(f"<b>{escape(c)}</b>", small) if c else "", Paragraph(escape(e), small) if e else ""]
          for (a, b), (c, e) in zip(meta[:half], meta[half:] + [("", "")] * (2 * half - len(meta)))]
    t = Table(kv, colWidths=[40 * mm, 51 * mm, 40 * mm, 51 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, hx("#DDDDDD")), ("BACKGROUND", (0, 0), (0, -1), hx("#F2F6FA")),
                           ("BACKGROUND", (2, 0), (2, -1), hx("#F2F6FA"))]))
    story += [t, Spacer(1, 2 * mm)]
    for n in _notes(rep):
        story.append(Paragraph(escape(n), ParagraphStyle("note", parent=small, textColor=hx("#8a5a00"))))
    def detail_block():
        det = rep.get("detalhe")
        if not det or not det["linhas"]:
            return []
        dcell = ParagraphStyle("dcell", parent=cell, fontSize=6.5, leading=8)
        dhead = ParagraphStyle("dhead", parent=dcell, textColor=colors.white, fontName="Helvetica-Bold")
        weights = {"usuario": 2.2, "destino": 2.4, "ips": 1.8, "maquinas": 1.6, "categoria": 1.5, "firewalls": 1.5,
                   "primeiro": 1.3, "ultimo": 1.3, "tipo": 0.9, "bloqueado_por": 1.6, "acessos": 1.05, "permitidos": 1.15, "bloqueados": 1.15}
        w = [weights.get(k, 0.9) for k, _ in det["colunas"]]
        widths = [doc.width * x / sum(w) for x in w]
        data = [[Paragraph(escape(label), dhead) for _, label in det["colunas"]]]
        data += [[Paragraph(escape(v[:90]), dcell) for v in r] for r in detail_rows(det)]
        t = LongTable(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), hx(HEAD)), ("GRID", (0, 0), (-1, -1), 0.25, hx("#BBBBBB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hx("#F2F6FA")]), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return [Paragraph(escape(det["titulo"]), h2), Paragraph(escape(detail_note(det)), small), Spacer(1, 1 * mm), t]

    if rep["tipo"] != "geral":  # quem acessou (ou o que o usuário acessou) vem logo depois do resumo
        story += detail_block()
    for chart in rep["graficos"]:
        fonte = f"Base: {_num(chart['base'])} registros ({chart['fonte']})."
        if chart.get("sem_valor"):
            fonte += f" {_num(chart['sem_valor'])} sem {chart['coluna'].lower()} no log."
        block = [Paragraph(escape(chart["titulo"]), h2), Paragraph(escape(fonte), small), Spacer(1, 1 * mm)]
        rows = _rows(chart)
        if chart["id"] == "acao":
            block.append(meter(chart))
        elif chart["tipo"] == "line":
            block.append(stacked(chart))
        else:
            graph = pie(chart) if chart["tipo"] == "pie" else hbar(chart)
            side = Table([[graph, table(_headers(chart), rows, [48 * mm, 22 * mm, 16 * mm])]],
                         colWidths=[92 * mm, 90 * mm])
            side.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
            block.append(side)
        story.append(KeepTogether(block))
    if rep["tipo"] == "geral":
        story += detail_block()
    if not rep["graficos"]:
        story.append(Paragraph("Nenhum evento no período para os filtros escolhidos.", st["Normal"]))
    doc.build(story)
    return out.getvalue()


# ---- Excel -----------------------------------------------------------------------------------
def to_xlsx(rep: dict, columns: list[tuple[str, str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, DoughnutChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    bold, head = Font(bold=True), PatternFill("solid", fgColor=HEAD[1:])
    ws.append([rep["titulo"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Gerado em {datetime.now():%d/%m/%Y %H:%M} pelo FortiLogPortal"])
    ws.append([])
    for k, v in _meta(rep):
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    for n in _notes(rep):
        ws.append([])
        ws.append([n])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    det = rep.get("detalhe")
    if det and det["linhas"]:
        sh = wb.create_sheet("Usuários" if det["agrupamento"] == "usuario" else "Destinos")
        sh.append([det["titulo"]])
        sh["A1"].font = Font(bold=True, size=12)
        sh.append([detail_note(det)])
        sh.append([label for _, label in det["colunas"]])
        for c in sh[3]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for line in det["linhas"]:
            sh.append([line.get(k) if isinstance(line.get(k), int) else _cell(line.get(k)) for k, _ in det["colunas"]])
        sh.freeze_panes = "A4"
        sh.auto_filter.ref = f"A3:{sh.cell(3, len(det['colunas'])).column_letter}{sh.max_row}"
        for i, (k, label) in enumerate(det["colunas"], start=1):
            sh.column_dimensions[sh.cell(3, i).column_letter].width = 30 if k in ("usuario", "destino", "ips", "maquinas") else max(12, len(label) + 3)
    used = {"Resumo", "Eventos", "Usuários", "Destinos"}
    for chart in rep["graficos"]:
        name = re.sub(r"[\[\]:*?/\\]", "", chart["titulo"])[:31]
        while name in used:
            name = name[:29] + "_" + str(len(used))
        used.add(name)
        sh = wb.create_sheet(name)
        sh.append([chart["titulo"]])
        sh["A1"].font = Font(bold=True, size=12)
        sh.append([f"Base: {chart['base']} registros ({chart['fonte']})"])
        sh.append(_headers(chart))
        for c in sh[3]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        if chart["tipo"] == "line":
            for k, a, b in chart["itens"]:
                sh.append([time_label(k, chart["agrupamento"]), a, b])
        else:
            total = sum(n for _, n in chart["itens"]) if chart["tipo"] == "pie" or chart["id"] == "acao" else chart["base"]
            for k, n in chart["itens"]:
                sh.append([_cell(k), n, round(n / total, 4) if total else 0])
                sh.cell(sh.max_row, 3).number_format = "0.0%"
        last = sh.max_row
        sh.column_dimensions["A"].width = 42
        sh.column_dimensions["B"].width = 14
        sh.column_dimensions["C"].width = 14
        if last < 4:
            continue
        cats = Reference(sh, min_col=1, min_row=4, max_row=last)
        if chart["tipo"] == "line":
            g = BarChart()
            g.type, g.grouping, g.overlap = "col", "stacked", 100
            g.add_data(Reference(sh, min_col=2, max_col=3, min_row=3, max_row=last), titles_from_data=True)
            g.set_categories(cats)
            for s, c in zip(g.series, (GOOD, CRITICAL)):
                s.graphicalProperties.solidFill = c[1:]
            g.width, g.height = 24, 9
        elif chart["tipo"] == "pie" or chart["id"] == "acao":
            g = DoughnutChart(holeSize=55)
            g.add_data(Reference(sh, min_col=2, min_row=3, max_row=last), titles_from_data=True)
            g.set_categories(cats)
            for i, c in enumerate(_colors(chart)):
                pt = DataPoint(idx=i)
                pt.graphicalProperties.solidFill = c[1:]
                g.series[0].dPt.append(pt)
            g.width, g.height = 14, 9
        else:
            g = BarChart()
            g.type = "bar"
            g.add_data(Reference(sh, min_col=2, min_row=3, max_row=last), titles_from_data=True)
            g.set_categories(cats)
            g.series[0].graphicalProperties.solidFill = SERIES[0][1:]
            g.y_axis.majorGridlines = None
            g.x_axis.scaling.orientation = "maxMin"  # o maior em cima, como na tela
            g.legend = None
            g.width, g.height = 16, max(7, 0.6 * (last - 3) + 2)
        g.title = chart["titulo"]
        sh.add_chart(g, "E3")

    ev = wb.create_sheet("Eventos")
    ev.append([label for _, label in columns])
    for c in ev[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for r in rep.get("eventos") or []:
        ev.append([_cell(r.get(k)) for k, _ in columns])
    ev.freeze_panes = "A2"
    for i, (_, label) in enumerate(columns, start=1):
        ev.column_dimensions[ev.cell(1, i).column_letter].width = max(12, len(label) + 4)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build(fmt: str, rep: dict, columns) -> tuple[bytes, str]:
    if fmt not in ("pdf", "xlsx"):
        raise ValueError("Formato deve ser pdf ou xlsx")
    content = to_pdf(rep) if fmt == "pdf" else to_xlsx(rep, [tuple(c) for c in columns])
    ascii_title = unicodedata.normalize("NFKD", rep["titulo"]).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9\-]+", "_", ascii_title.lower()).strip("_")[:50] or "relatorio"
    filename = f"{slug}_{datetime.now():%Y%m%d_%H%M%S}.{fmt}"
    settings.export_dir.mkdir(parents=True, exist_ok=True)
    Path(settings.export_dir / filename).write_bytes(content)
    return content, filename


__all__ = ["FORMATS", "build", "time_label"]
